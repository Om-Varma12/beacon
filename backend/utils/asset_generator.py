import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE

def set_cell_background(cell, color_hex):
    """Sets the background color of a table cell in a DOCX document."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_docx_proposal(doc_data: dict) -> io.BytesIO:
    """Generates a beautifully styled Word Document (.docx) proposal from the assets JSON."""
    doc = Document()
    
    # Page Setup (Margins: 1 inch on all sides)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Set default fonts and style formatting
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(45, 49, 66)  # Charcoal
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)
    
    # --- Title Page Header ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(12)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(doc_data.get("title", "Technical Proposal").upper())
    title_run.font.name = 'Segoe UI'
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(27, 67, 50)  # Forest Green
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(24)
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("AI-Generated Hackathon Solution Blueprint")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(64, 145, 108)  # Mint Green
    
    # Separator Line
    doc.add_paragraph("━" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Helper functions to add headers with consistent typography
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(27, 67, 50)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(64, 145, 108)
        return p

    # --- Section 1: Executive Summary ---
    add_heading_1("1. Solution Overview")
    doc.add_paragraph(doc_data.get("solution_overview", ""))
    
    # --- Section 2: Problem Statement & Rationale ---
    add_heading_1("2. Problem Statement")
    doc.add_paragraph(doc_data.get("problem_statement", ""))
    
    if doc_data.get("why_this_is_best"):
        add_heading_2("Why This Concept Wins")
        doc.add_paragraph(doc_data.get("why_this_is_best", ""))
        
    # --- Section 3: Core Features & User Personas ---
    add_heading_1("3. Core Features & User Personas")
    
    add_heading_2("Core Features")
    for feat in doc_data.get("core_features", []):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(feat)
        
    add_heading_2("Target Users / Personas")
    for user in doc_data.get("target_users", []):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(user)
        
    # --- Section 4: Recommended Tech Stack Table ---
    tech_stack = doc_data.get("recommended_tech_stack", {})
    if tech_stack:
        add_heading_1("4. Recommended Technical Architecture")
        
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Light Shading Accent 1'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Architecture Layer'
        hdr_cells[1].text = 'Technologies'
        
        set_cell_background(hdr_cells[0], "1B4332")
        set_cell_background(hdr_cells[1], "1B4332")
        
        for cell in hdr_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                    r.font.size = Pt(10)
        
        layers_map = [
            ("Frontend / Client", tech_stack.get("frontend", [])),
            ("Backend API / Logic", tech_stack.get("backend", [])),
            ("Database & Caching", tech_stack.get("database", [])),
            ("AI & Machine Learning", tech_stack.get("ml_ai", [])),
            ("DevOps & Hosting", tech_stack.get("devops", [])),
            ("Third-Party Integrations", tech_stack.get("apis_services", []))
        ]
        
        for layer_name, techs in layers_map:
            if techs:
                row_cells = table.add_row().cells
                row_cells[0].text = layer_name
                row_cells[1].text = ", ".join(techs)
                
                for cell in row_cells:
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        for r in p.runs:
                            r.font.size = Pt(9.5)
        doc.add_paragraph()

    # --- Section 5: Implementation Roadmap ---
    phases = doc_data.get("implementation_phases", [])
    if phases:
        add_heading_1("5. Implementation Roadmap")
        for p_data in phases:
            phase_name = p_data.get("phase", "Phase")
            add_heading_2(phase_name)
            for task in p_data.get("tasks", []):
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(task)
                
    # --- Section 6: Team Structure ---
    team = doc_data.get("team_structure", [])
    if team:
        add_heading_1("6. Team Structure & Roles")
        for member in team:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(member)
            
    # --- Section 7: Unique Selling Points ---
    usps = doc_data.get("unique_selling_points", [])
    if usps:
        add_heading_1("7. Business Value & USPs")
        for usp in usps:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(usp)
            
    # --- Section 8: Risks & Mitigations Table ---
    risks = doc_data.get("risks_and_mitigations", [])
    if risks:
        add_heading_1("8. Risk Management & Mitigations")
        
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Light Shading Accent 1'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Risk Identified'
        hdr_cells[1].text = 'Mitigation Strategy'
        
        set_cell_background(hdr_cells[0], "1B4332")
        set_cell_background(hdr_cells[1], "1B4332")
        
        for cell in hdr_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                    r.font.size = Pt(10)
                    
        for r_data in risks:
            row_cells = table.add_row().cells
            row_cells[0].text = r_data.get("risk", "")
            row_cells[1].text = r_data.get("mitigation", "")
            for cell in row_cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
        doc.add_paragraph()

    # --- Section 9: Judging Criteria Alignment ---
    judging = doc_data.get("judging_criteria_alignment", [])
    if judging:
        add_heading_1("9. Judging Criteria Alignment")
        for item in judging:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(item)
            
    # Save the generated document into the memory stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

def generate_pptx_slides(ppt_data: dict) -> io.BytesIO:
    """Generates a beautifully styled PowerPoint Presentation (.pptx) deck from the assets JSON."""
    prs = Presentation()
    
    # 16:9 Standard widescreen layout
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    title_text = ppt_data.get("title", "Project Proposal").upper()
    organizer_text = ppt_data.get("organizer", "Hackathon Organizer")
    
    def add_textbox(slide, left, top, width, height, margin=0):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = PptInches(margin)
        tf.margin_right = PptInches(margin)
        tf.margin_top = PptInches(margin)
        tf.margin_bottom = PptInches(margin)
        return tf
        
    def set_slide_background(slide, rgb_tuple):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = PptRGBColor(*rgb_tuple)
        
    # --- SLIDE 1: Title Slide (Dark Theme) ---
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, (27, 67, 50))  # Forest Green
    
    tf_title = add_textbox(slide1, PptInches(1.0), PptInches(2.2), PptInches(11.333), PptInches(4.0))
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = 'Segoe UI'
    p_title.font.size = PptPt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = PptRGBColor(255, 255, 255)
    
    p_sub = tf_title.add_paragraph()
    p_sub.text = f"HACKATHON SOLUTION BLUEPRINT\n{organizer_text.upper()}"
    p_sub.font.name = 'Segoe UI'
    p_sub.font.size = PptPt(16)
    p_sub.font.bold = False
    p_sub.font.color.rgb = PptRGBColor(64, 145, 108)  # Mint Green
    p_sub.space_before = PptPt(24)
    
    # --- CONTENT SLIDES (Light Theme) ---
    slides_data = ppt_data.get("slides", [])
    for slide_info in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide, (248, 249, 250))  # Light Grey
        
        # Horizontal forest green Accent line at the very top of each content slide
        accent_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            PptInches(0), PptInches(0), PptInches(13.333), PptInches(0.15)
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = PptRGBColor(27, 67, 50)
        accent_bar.line.fill.background()
        
        # Slide Header
        slide_title = slide_info.get("title", "Slide Title").upper()
        tf_header = add_textbox(slide, PptInches(1.0), PptInches(0.6), PptInches(11.333), PptInches(1.0))
        p_hdr = tf_header.paragraphs[0]
        p_hdr.text = slide_title
        p_hdr.font.name = 'Segoe UI'
        p_hdr.font.size = PptPt(28)
        p_hdr.font.bold = True
        p_hdr.font.color.rgb = PptRGBColor(27, 67, 50)
        
        # Slide Body (Bullet Points)
        bullets = slide_info.get("bullets", [])
        tf_body = add_textbox(slide, PptInches(1.0), PptInches(1.8), PptInches(11.333), PptInches(4.5))
        
        for idx, bullet in enumerate(bullets):
            p = tf_body.paragraphs[0] if idx == 0 else tf_body.add_paragraph()
            p.text = "•  " + bullet
            p.font.name = 'Segoe UI'
            p.font.size = PptPt(16)
            p.font.color.rgb = PptRGBColor(45, 49, 66)  # Charcoal
            p.space_after = PptPt(14)
            p.level = 0
            
        # Slide Footer
        tf_footer = add_textbox(slide, PptInches(1.0), PptInches(6.8), PptInches(11.333), PptInches(0.5))
        p_ft = tf_footer.paragraphs[0]
        p_ft.text = f"{title_text}  |  Slide {slide_info.get('slide_number', '')}"
        p_ft.font.name = 'Segoe UI'
        p_ft.font.size = PptPt(9)
        p_ft.font.color.rgb = PptRGBColor(160, 160, 160)
        
        # Add Speaker Notes to slide
        notes_slide = slide.notes_slide
        notes_text_frame = notes_slide.notes_text_frame
        notes_text_frame.text = slide_info.get("speaker_notes", "")
        
    # Save the generated presentation into the memory stream
    file_stream = io.BytesIO()
    prs.save(file_stream)
    file_stream.seek(0)
    return file_stream
