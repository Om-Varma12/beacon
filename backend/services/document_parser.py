import io
import os
import mimetypes
import httpx
import pdfplumber

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    from pptx import Presentation as PptxPresentation
except Exception:
    PptxPresentation = None

from bs4 import BeautifulSoup


class DocumentParser:
    """Simple document parser for PDFs, DOCX, TXT, PPTX and basic HTML.

    Downloads a resource URL and returns a dictionary with parsed text
    and basic metadata.
    """

    async def fetch(self, url: str) -> dict:
        async with httpx.AsyncClient(timeout=30.0, follow_redirects=True) as client:
            resp = await client.get(url)
            resp.raise_for_status()
            content_type = resp.headers.get("content-type", "")
            return {"content": resp.content, "content_type": content_type}

    def _detect_type(self, url: str, content_type: str) -> str:
        # Prefer extension
        path = url.split("?")[0]
        ext = os.path.splitext(path)[1].lower()
        if ext in (".pdf",):
            return "pdf"
        if ext in (".docx", ".doc"):
            return "docx"
        if ext in (".pptx", ".ppt"):
            return "pptx"
        if ext in (".txt",):
            return "txt"
        # fallback to content-type
        if "pdf" in content_type:
            return "pdf"
        if "word" in content_type or "officedocument" in content_type:
            return "docx"
        if "presentation" in content_type:
            return "pptx"
        if "text" in content_type:
            return "txt"
        if "html" in content_type:
            return "html"

        return "binary"

    def _parse_pdf(self, content: bytes) -> str:
        text_chunks = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_chunks.append(text)
        return "\n\n".join(text_chunks)

    def _parse_docx(self, content: bytes) -> str:
        if DocxDocument is None:
            return ""
        doc = DocxDocument(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        return "\n\n".join(paragraphs)

    def _parse_pptx(self, content: bytes) -> str:
        if PptxPresentation is None:
            return ""
        prs = PptxPresentation(io.BytesIO(content))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                try:
                    if hasattr(shape, "text"):
                        t = shape.text
                        if t:
                            texts.append(t)
                except Exception:
                    continue
        return "\n\n".join(texts)

    def _parse_html(self, content: bytes) -> str:
        soup = BeautifulSoup(content, "html.parser")
        # Get visible text
        for s in soup(['script', 'style', 'noscript']):
            s.extract()
        text = soup.get_text(separator="\n")
        return text

    def _parse_txt(self, content: bytes) -> str:
        try:
            return content.decode("utf-8")
        except Exception:
            try:
                return content.decode("latin-1")
            except Exception:
                return ""

    async def parse(self, url: str) -> dict:
        fetched = await self.fetch(url)
        content = fetched.get("content", b"")
        content_type = fetched.get("content_type", "")
        doc_type = self._detect_type(url, content_type)

        parsers = {
            "pdf": self._parse_pdf,
            "docx": self._parse_docx,
            "pptx": self._parse_pptx,
            "html": self._parse_html,
            "txt": self._parse_txt,
        }

        parser = parsers.get(doc_type)
        text = ""
        if parser:
            try:
                text = parser(content)
            except Exception:
                text = ""

        return {
            "source_url": url,
            "source_type": doc_type,
            "text": text,
            "confidence": 0.9 if text else 0.1,
        }
